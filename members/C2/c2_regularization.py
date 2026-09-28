"""
C2 담당 - 드롭아웃 / 조기종료 / 클래스 가중치(정규화·불균형 대응) 실험
- 9/28~9/29: 조기종료 기본 구현 점검 (monitor, patience 비교)
- 9/29~10/1: 드롭아웃 비율 실험
- 10/1~10/3: 클래스 가중치 실험
- 10/5~ : C1의 상위 구조와 조합해 재검증
끝내다


전제 (C1과 동일하게 맞춤 - 비교 공정성 확보):
- A가 나눠준 train.csv / test.csv 사용 (train_test_split 80:20, stratify=Revenue, seed=42)
- B의 평가 스크립트 아직 없어서, train.csv를 다시 train/val로 "임시" 분할해서 사용
  -> B 스크립트 오면 3번 섹션만 교체할 것
- 전처리도 "임시": Weekend를 int로, Month/VisitorType은 원-핫 (C1과 동일한 함수)
  -> A의 공통 전처리 오면 이 부분 교체할 것
- 기반 구조(은닉층 수·노드 수)는 C1이 확정한 최상위 설정을 그대로 가져다 쓴다.
  지금은 C1의 baseline 구조 (64, 32)를 임시로 사용 중 -> C1 구조 실험 끝나면 HIDDEN_UNITS만 교체
- 같은 폴더의 experiment_log_C2.xlsx(팀 공통 양식, C2용)에 결과를 자동으로 한 줄씩 추가함

실험 1회 = 이 스크립트 1번 실행. STAGE / DROPOUT_RATE / EARLY_STOP_* / USE_CLASS_WEIGHT
값을 바꿔가며 여러 번 돌리고, 매번 로그에 한 줄씩 쌓인다.
"""

import os
from datetime import date

import numpy as np
import pandas as pd
import tensorflow as tf
from tensorflow import keras
from tensorflow.keras import layers
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.utils.class_weight import compute_class_weight
from sklearn.metrics import recall_score, f1_score, average_precision_score, classification_report
import openpyxl

# ---- 경로 설정 (이 스크립트와 같은 폴더 기준) ----
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "..", "..", "data")
LOG_PATH = os.path.join(BASE_DIR, "experiment_log_C2.xlsx")  # 팀 공통 엑셀 양식(C2용)

# ============================================================
# ===== 이 블록만 바꿔가며 실험을 반복한다 =====
# ============================================================
STAGE = "dropout"              # "dropout" / "early_stop" / "class_weight" / "combined" / "final"
HIDDEN_UNITS = (64, 32)        # C1 확정 구조로 교체 예정 (지금은 C1 baseline 구조)
DROPOUT_RATE = 0.3             # 0.0이면 드롭아웃 미적용. 층마다 다르게 하려면 튜플로: (0.3, 0.2)
USE_EARLY_STOP = True
EARLY_STOP_MONITOR = "val_loss"   # 불균형 데이터라 val_f1/val_pr_auc 커스텀도 고려 중
EARLY_STOP_PATIENCE = 5
USE_CLASS_WEIGHT = False
THRESHOLD = 0.5
MAX_EPOCHS = 50
BATCH_SIZE = 32
SEED = 42
MEMO = "드롭아웃 0.3, 나머지는 C1 baseline과 동일 조건"
# ============================================================

np.random.seed(SEED)
tf.random.set_seed(SEED)

# ---- 1. 데이터 불러오기 (A가 나눠준 train/test, C1과 동일) ----
train_df = pd.read_csv(os.path.join(DATA_DIR, "train.csv"))
test_df = pd.read_csv(os.path.join(DATA_DIR, "test.csv"))
print("train:", train_df.shape, "test:", test_df.shape)


# ---- 2. 임시 전처리 (C1과 동일한 함수 - 팀 공통 전처리 오면 통째로 교체) ----
PREP_VERSION = "temp_v1(원-핫, A공통전처리 전)"


def preprocess(df, ref_columns=None):
    df = df.copy()
    df["Weekend"] = df["Weekend"].astype(int)
    y = df["Revenue"].astype(int) if "Revenue" in df.columns else None
    X = df.drop(columns=["Revenue"]) if "Revenue" in df.columns else df
    X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=True)
    if ref_columns is not None:
        X = X.reindex(columns=ref_columns, fill_value=0)
    return X, y


X_train_full, y_train_full = preprocess(train_df)
X_test, y_test = preprocess(test_df, ref_columns=X_train_full.columns)

# ---- 3. 임시 train/val 분할 (C1과 동일 - B 스크립트 오면 교체) ----
X_train, X_val, y_train, y_val = train_test_split(
    X_train_full, y_train_full,
    test_size=0.2, random_state=SEED, stratify=y_train_full
)

scaler = StandardScaler()
X_train_s = scaler.fit_transform(X_train)
X_val_s = scaler.transform(X_val)
X_test_s = scaler.transform(X_test)
print("train/val/test shape:", X_train_s.shape, X_val_s.shape, X_test_s.shape)


# ---- 4. MLP (C1 구조 + 드롭아웃) ----
def build_model(input_dim, hidden_units=HIDDEN_UNITS, dropout_rate=DROPOUT_RATE):
    if isinstance(dropout_rate, (int, float)):
        dropout_rate = [dropout_rate] * len(hidden_units)

    model = keras.Sequential([layers.Input(shape=(input_dim,))])
    for units, drop in zip(hidden_units, dropout_rate):
        model.add(layers.Dense(units, activation="relu"))
        if drop and drop > 0:
            model.add(layers.Dropout(drop))
    model.add(layers.Dense(1, activation="sigmoid"))

    model.compile(optimizer="adam", loss="binary_crossentropy", metrics=["accuracy"])
    return model


model = build_model(X_train_s.shape[1])
model.summary()

# ---- 5. 클래스 가중치 계산 (불균형 15.5% 대응) ----
if USE_CLASS_WEIGHT:
    classes = np.unique(y_train)
    weights = compute_class_weight(class_weight="balanced", classes=classes, y=y_train)
    class_weight_dict = {int(c): float(w) for c, w in zip(classes, weights)}
else:
    class_weight_dict = None

print("class_weight:", class_weight_dict)

# ---- 6. 조기종료 콜백 ----
callbacks = []
if USE_EARLY_STOP:
    callbacks.append(
        keras.callbacks.EarlyStopping(
            monitor=EARLY_STOP_MONITOR,
            patience=EARLY_STOP_PATIENCE,
            restore_best_weights=True,
        )
    )

# ---- 7. 학습 ----
history = model.fit(
    X_train_s, y_train,
    validation_data=(X_val_s, y_val),
    epochs=MAX_EPOCHS,
    batch_size=BATCH_SIZE,
    class_weight=class_weight_dict,
    callbacks=callbacks,
    verbose=1,
)
actual_epochs = len(history.history["loss"])

# ---- 8. 평가: 정확도(X) 대신 재현율/F1/PR-AUC ----
y_val_pred_prob = model.predict(X_val_s).ravel()
y_val_pred = (y_val_pred_prob >= THRESHOLD).astype(int)

recall = recall_score(y_val, y_val_pred)
f1 = f1_score(y_val, y_val_pred)
pr_auc = average_precision_score(y_val, y_val_pred_prob)

print(f"\n[검증셋] Recall: {recall:.4f}  F1: {f1:.4f}  PR-AUC: {pr_auc:.4f}  (실제 학습 에폭: {actual_epochs})")
print(classification_report(y_val, y_val_pred, target_names=["No Purchase", "Purchase"]))


# ---- 9. 실험 로그 엑셀에 자동 기록 ----
def log_to_excel(log_path, entry, sheet_name="실험로그"):
    """팀 공통 엑셀 양식(실험로그 시트) 맨 아래 빈 줄에 한 행 추가."""
    wb = openpyxl.load_workbook(log_path)
    ws = wb[sheet_name]

    # A열이 비어있는 첫 행 찾기 (헤더 1행 + 예시 2행 다음부터)
    next_row = 3
    while ws.cell(row=next_row, column=1).value is not None:
        next_row += 1

    exp_id = f"C2-{next_row - 2:03d}"

    row_values = [
        exp_id,
        entry["date"],
        entry["stage"],
        entry["framework"],
        entry["prep_version"],
        entry["input_dim"],
        entry["hidden_layers"],
        entry["nodes"],
        entry["dropout_rate"],
        entry["es_monitor"],
        entry["es_patience"],
        entry["class_weight"],
        entry["threshold"],
        entry["max_epochs"],
        entry["actual_epochs"],
        entry["seed"],
        entry["recall"],
        entry["f1"],
        entry["pr_auc"],
        entry["memo"],
    ]
    for col, value in enumerate(row_values, start=1):
        ws.cell(row=next_row, column=col, value=value)

    wb.save(log_path)
    print(f"[로그 저장 완료] {exp_id} -> {log_path}")
    return exp_id


log_entry = {
    "date": date.today().isoformat(),
    "stage": STAGE,
    "framework": f"tensorflow(keras) {tf.__version__}",
    "prep_version": PREP_VERSION,
    "input_dim": X_train_s.shape[1],
    "hidden_layers": len(HIDDEN_UNITS),
    "nodes": str(HIDDEN_UNITS),
    "dropout_rate": str(DROPOUT_RATE),
    "es_monitor": EARLY_STOP_MONITOR if USE_EARLY_STOP else "없음",
    "es_patience": EARLY_STOP_PATIENCE if USE_EARLY_STOP else "없음",
    "class_weight": str(class_weight_dict) if class_weight_dict else "없음",
    "threshold": THRESHOLD,
    "max_epochs": MAX_EPOCHS,
    "actual_epochs": actual_epochs,
    "seed": SEED,
    "recall": round(float(recall), 4),
    "f1": round(float(f1), 4),
    "pr_auc": round(float(pr_auc), 4),
    "memo": MEMO,
}

log_to_excel(LOG_PATH, log_entry)

