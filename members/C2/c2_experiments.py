"""
C2 담당 - 드롭아웃 / 조기종료 / 클래스 가중치 실험 (C1 기준 버전)

사용법:
    python c2_experiments.py dropout        # 드롭아웃 비율 실험
    python c2_experiments.py early_stop     # 조기종료 monitor / patience 실험
    python c2_experiments.py class_weight   # 클래스 가중치 실험
    python c2_experiments.py combined       # 조합 재검증 (10/5~)
    python c2_experiments.py all            # dropout + early_stop + class_weight (기본값)

C1 기준으로 맞춘 점 (비교 공정성):
- 데이터·전처리·train/val 분할·스케일링: c1_experiments.py와 완전히 동일
- BASE 구조·옵티마이저·학습률·배치 크기 = C1이 확정한 최상위 설정 (아래 BASE만 교체)
- 한 번에 하나만 바꾸는 one-factor-at-a-time, 여러 시드 평균 ± 표준편차
- 비교 대상(기준선)은 "정규화 없음" 조건: 드롭아웃 0, 조기종료 없음, 클래스 가중치 없음, 고정 50 epoch
  (= C1 baseline과 같은 조건이므로 C1 결과와 직접 비교 가능)
- 결과는 experiment_log_C2.xlsx에 한 줄씩 자동 기록 + results_C2_<stage>.csv 요약 저장

주의: 조기종료는 val 셋으로 멈추는 시점을 고르기 때문에, 같은 val로 평가하면 약간 낙관적으로 나온다.
      최종 모델(10/14 동결) 전에는 별도 test 평가(B 스크립트)로 재확인할 것.
"""

import os
import sys
from datetime import date

import numpy as np
import pandas as pd
import tensorflow as tf
from tensorflow import keras
from tensorflow.keras import layers
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.utils.class_weight import compute_class_weight
from sklearn.metrics import recall_score, f1_score, average_precision_score
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "..", "..", "data")
LOG_PATH = os.path.join(BASE_DIR, "experiment_log_C2.xlsx")

# ============================================================
# ===== 실험 설정 =====
# ============================================================
SPLIT_SEED = 42                 # 데이터 분할 시드 (고정, C1과 동일)
SEEDS = [42, 1, 2]              # 가중치 초기화 시드 (C1과 동일)
MAX_EPOCHS = 50
THRESHOLD = 0.5

# C1이 확정한 최상위 설정으로 교체 (지금은 C1 baseline)
BASE = dict(
    hidden=(64, 32), optimizer="adam", lr=1e-3, batch=32,          # <- C1 확정값
    dropout=0.0, es_monitor=None, es_patience=None, class_weight=False,  # <- C2가 바꾸는 부분
)

GRID = {
    "dropout": [dict(dropout=d) for d in [0.0, 0.1, 0.2, 0.3, 0.5]],
    "early_stop": [
        dict(es_monitor="val_loss", es_patience=3),
        dict(es_monitor="val_loss", es_patience=5),
        dict(es_monitor="val_loss", es_patience=10),
        dict(es_monitor="val_pr_auc", es_patience=5),   # 불균형 데이터용 (PR-AUC 기준)
        dict(es_monitor="val_pr_auc", es_patience=10),
    ],
    "class_weight": [dict(class_weight=False), dict(class_weight=True)],
    # 개별 실험에서 좋았던 값으로 채워 넣을 것 (기본값은 예시)
    "combined": [
        dict(dropout=0.3, es_monitor="val_loss", es_patience=5, class_weight=False),
        dict(dropout=0.3, es_monitor="val_loss", es_patience=5, class_weight=True),
        dict(dropout=0.2, es_monitor="val_pr_auc", es_patience=5, class_weight=True),
    ],
}
# ============================================================


# ---- 1. 데이터 (C1과 동일) ----
train_df = pd.read_csv(os.path.join(DATA_DIR, "train.csv"))

PREP_VERSION = "temp_v1(원-핫, A공통전처리 전)"


def preprocess(df, ref_columns=None):
    df = df.copy()
    df["Weekend"] = df["Weekend"].astype(int)
    y = df["Revenue"].astype(int) if "Revenue" in df.columns else None
    X = df.drop(columns=["Revenue"]) if "Revenue" in df.columns else df
    X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=True)
    if ref_columns is not None:
        X = X.reindex(columns=ref_columns, fill_value=0)
    return X, y


X_full, y_full = preprocess(train_df)

# ---- 2. 임시 train/val 분할 (C1과 동일 - B 스크립트 오면 교체) ----
X_train, X_val, y_train, y_val = train_test_split(
    X_full, y_full, test_size=0.2, random_state=SPLIT_SEED, stratify=y_full
)
scaler = StandardScaler()
X_train_s = scaler.fit_transform(X_train)
X_val_s = scaler.transform(X_val)
INPUT_DIM = X_train_s.shape[1]


# ---- 3. 모델 (C1 구조 + 드롭아웃) ----
def make_optimizer(name, lr):
    if name == "adam":
        return keras.optimizers.Adam(learning_rate=lr)
    if name == "rmsprop":
        return keras.optimizers.RMSprop(learning_rate=lr)
    if name == "sgd":
        return keras.optimizers.SGD(learning_rate=lr, momentum=0.9)
    raise ValueError(f"지원하지 않는 옵티마이저: {name}")


def build_model(cfg):
    drop = cfg["dropout"]
    drops = [drop] * len(cfg["hidden"]) if isinstance(drop, (int, float)) else list(drop)

    model = keras.Sequential([layers.Input(shape=(INPUT_DIM,))])
    for units, d in zip(cfg["hidden"], drops):
        model.add(layers.Dense(units, activation="relu"))
        if d and d > 0:
            model.add(layers.Dropout(d))
    model.add(layers.Dense(1, activation="sigmoid"))

    model.compile(
        optimizer=make_optimizer(cfg["optimizer"], cfg["lr"]),
        loss="binary_crossentropy",
        metrics=["accuracy", keras.metrics.AUC(curve="PR", name="pr_auc")],  # val_pr_auc 모니터용
    )
    return model


def get_class_weight(use):
    if not use:
        return None
    classes = np.unique(y_train)
    w = compute_class_weight("balanced", classes=classes, y=y_train)
    return {int(c): float(v) for c, v in zip(classes, w)}


def run_once(cfg, seed):
    np.random.seed(seed)
    tf.random.set_seed(seed)
    keras.backend.clear_session()

    model = build_model(cfg)

    callbacks = []
    if cfg["es_monitor"]:
        callbacks.append(keras.callbacks.EarlyStopping(
            monitor=cfg["es_monitor"],
            mode="max" if cfg["es_monitor"] == "val_pr_auc" else "min",
            patience=cfg["es_patience"],
            restore_best_weights=True,
        ))

    hist = model.fit(
        X_train_s, y_train, validation_data=(X_val_s, y_val),
        epochs=MAX_EPOCHS, batch_size=cfg["batch"],
        class_weight=get_class_weight(cfg["class_weight"]),
        callbacks=callbacks, verbose=0,
    )

    prob = model.predict(X_val_s, verbose=0).ravel()
    pred = (prob >= THRESHOLD).astype(int)
    return dict(
        recall=recall_score(y_val, pred),
        f1=f1_score(y_val, pred),
        pr_auc=average_precision_score(y_val, prob),
        actual_epochs=len(hist.history["loss"]),
    )


# ---- 4. 엑셀 로그 (C2 양식: 헤더 1행 + 예시 2행 다음부터) ----
def log_to_excel(entry, sheet_name="실험로그"):
    wb = openpyxl.load_workbook(LOG_PATH)
    ws = wb[sheet_name]
    next_row = 3
    while ws.cell(row=next_row, column=1).value is not None:
        next_row += 1
    exp_id = f"C2-{next_row - 2:03d}"

    values = [
        exp_id, entry["date"], entry["stage"], entry["framework"], entry["prep_version"],
        entry["input_dim"], entry["hidden_layers"], entry["nodes"], entry["dropout_rate"],
        entry["es_monitor"], entry["es_patience"], entry["class_weight"], entry["threshold"],
        entry["max_epochs"], entry["actual_epochs"], entry["seed"],
        entry["recall"], entry["f1"], entry["pr_auc"], entry["memo"],
    ]
    for col, v in enumerate(values, start=1):
        ws.cell(row=next_row, column=col, value=v)
    wb.save(LOG_PATH)
    return exp_id


def run_experiment(stage, changed, cfg):
    df = pd.DataFrame([run_once(cfg, s) for s in SEEDS])
    mean, std = df.mean(), df.std(ddof=0)

    cw = get_class_weight(cfg["class_weight"])
    memo = (f"{stage}: {changed} | 기반 {cfg['hidden']}, {cfg['optimizer']}, "
            f"lr={cfg['lr']}, batch={cfg['batch']} | seeds={SEEDS} 평균 | "
            f"F1 std={std['f1']:.4f}, Recall std={std['recall']:.4f}")

    exp_id = log_to_excel(dict(
        date=date.today().isoformat(), stage=stage,
        framework=f"tensorflow(keras) {tf.__version__}", prep_version=PREP_VERSION,
        input_dim=INPUT_DIM, hidden_layers=len(cfg["hidden"]), nodes=str(cfg["hidden"]),
        dropout_rate=str(cfg["dropout"]),
        es_monitor=cfg["es_monitor"] or "없음",
        es_patience=cfg["es_patience"] if cfg["es_monitor"] else "없음",
        class_weight=str({k: round(v, 3) for k, v in cw.items()}) if cw else "없음",
        threshold=THRESHOLD, max_epochs=MAX_EPOCHS,
        actual_epochs=round(float(mean["actual_epochs"]), 1),
        seed=f"{SEEDS} 평균",
        recall=round(float(mean["recall"]), 4), f1=round(float(mean["f1"]), 4),
        pr_auc=round(float(mean["pr_auc"]), 4), memo=memo,
    ))
    print(f"[{exp_id}] {changed:<55} Recall {mean['recall']:.4f}  "
          f"F1 {mean['f1']:.4f}±{std['f1']:.4f}  PR-AUC {mean['pr_auc']:.4f}  "
          f"epochs {mean['actual_epochs']:.1f}")
    return dict(exp_id=exp_id, stage=stage, **{k: str(v) for k, v in changed_dict(cfg).items()},
                recall=mean["recall"], f1=mean["f1"], f1_std=std["f1"],
                pr_auc=mean["pr_auc"], actual_epochs=mean["actual_epochs"])


def changed_dict(cfg):
    return {k: cfg[k] for k in ("dropout", "es_monitor", "es_patience", "class_weight")}


# ---- 5. 실행 ----
if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "all"
    stages = ["dropout", "early_stop", "class_weight"] if target == "all" else [target]

    for stage in stages:
        print(f"\n===== {stage} 실험 시작 =====")
        rows = []
        for change in GRID[stage]:
            cfg = {**BASE, **change}
            rows.append(run_experiment(stage, str(change), cfg))
        out = os.path.join(BASE_DIR, f"results_C2_{stage}.csv")
        pd.DataFrame(rows).sort_values("pr_auc", ascending=False).to_csv(
            out, index=False, encoding="utf-8-sig")
        print(f"[요약 저장] {out}")
