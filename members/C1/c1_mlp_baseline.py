"""
C1 담당 - MLP 기본 골격 + 학습/평가 루프 + 실험 로그 자동 기록
- 9/22: MLP(입력->은닉2층->출력) + 학습/평가 루프
- 9/23: 첫 학습 + 재현율/F1 측정, 로그 기록 시작
- 9/27: 엑셀 로그 자동 기록 추가
- 9/29: 공식 분할(train_clean / validation_final)만 사용하도록 변경

전제:
- A의 공식 분할을 data/final/train_clean.csv, data/final/validation_final.csv로 사용
  (make_final_split.py로 복사한 파일. train을 다시 나누지 않음)
- 보류 분할은 이 스크립트에서 읽지 않음. 모델 확정 후 별도 평가에서만 사용
- 전처리는 아직 임시: Weekend를 int로, Month/VisitorType은 원-핫
  -> A의 공통 전처리 오면 이 부분 교체할 것
- 같은 폴더에 있는 experiment_log_C1.xlsx(팀 공통 양식)에 결과를 자동으로 한 줄씩 추가함
"""

import argparse
import os
import random
from datetime import date

import numpy as np
import pandas as pd
import tensorflow as tf
from tensorflow import keras
from tensorflow.keras import layers
from sklearn.preprocessing import StandardScaler
from sklearn.utils.class_weight import compute_class_weight
from sklearn.metrics import (
    recall_score,
    f1_score,
    average_precision_score,
    classification_report,
    precision_score,
    confusion_matrix,
)
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument("--seed", type=int, default=42)
args = parser.parse_args()
SEED = args.seed

# ---- 실험 조건 (임시값. TODO: C2와 합의 후 확정) ----
HIDDEN_UNITS = (64, 32)  # 은닉층 노드 수 (입력 다음). TODO: C2와 합의 후 확정
LEARNING_RATE = 1e-3  # Adam 학습률. TODO: C2와 합의 후 확정
OPTIMIZER = keras.optimizers.Adam(learning_rate=LEARNING_RATE)  # 실제 Keras 옵티마이저. TODO: C2와 합의 후 확정
BATCH_SIZE = 32  # 미니배치 크기. TODO: C2와 합의 후 확정
MAX_EPOCHS = 50  # 최대 에폭 (조기종료가 더 일찍 끝낼 수 있음). TODO: C2와 합의 후 확정
EARLY_STOPPING = True  # 조기종료 사용 여부. TODO: C2와 합의 후 확정
ES_MONITOR = "val_pr_auc"  # 조기종료가 보는 지표. TODO: C2와 합의 후 확정
ES_MODE = "max"  # 지표가 커질수록 좋다고 봄. TODO: C2와 합의 후 확정
ES_PATIENCE = 5  # 개선이 없어도 기다리는 에폭 수. TODO: C2와 합의 후 확정
RESTORE_BEST_WEIGHTS = True  # 가장 좋았던 가중치로 되돌림. TODO: C2와 합의 후 확정
USE_CLASS_WEIGHT = True  # 구매/비구매 비율로 class_weight 사용. TODO: C2와 합의 후 확정

# ---- 경로 설정 (이 스크립트와 같은 폴더 기준) ----
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "..", "..", "data")
FINAL_DIR = os.path.join(DATA_DIR, "final")
LOG_PATH = os.path.join(BASE_DIR, "experiment_log_C1.xlsx")  # 팀 공통 엑셀 양식

# ---- 0. 시드 고정 ----
random.seed(SEED)
np.random.seed(SEED)
tf.random.set_seed(SEED)

# ---- 1. 데이터 불러오기 (공식 분할) ----
train_df = pd.read_csv(os.path.join(FINAL_DIR, "train_clean.csv"))
val_df = pd.read_csv(os.path.join(FINAL_DIR, "validation_final.csv"))
print("train:", train_df.shape, "val:", val_df.shape)

for col in ("Month", "VisitorType"):
    extra = set(val_df[col].astype(str)) - set(train_df[col].astype(str))
    print(f"[val only {col}]", sorted(extra) if extra else "없음")

# ---- 2. 임시 전처리 (팀 공통 전처리 오면 이 함수 통째로 교체) ----
PREP_VERSION = "temp_v1(원-핫, A공통전처리 전)"


def preprocess(df, ref_columns=None):
    df = df.copy()
    df["Weekend"] = df["Weekend"].astype(int)
    y = df["Revenue"].astype(int) if "Revenue" in df.columns else None
    X = df.drop(columns=["Revenue"]) if "Revenue" in df.columns else df
    if ref_columns is None:
        X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=True)
    else:
        X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=False)
        X = X.reindex(columns=ref_columns, fill_value=0)
    return X, y


X_train, y_train = preprocess(train_df)
X_val, y_val = preprocess(val_df, ref_columns=X_train.columns)
X_val = X_val[X_train.columns]
assert list(X_train.columns) == list(X_val.columns)
print("train/val columns equal:", list(X_train.columns) == list(X_val.columns))

scaler = StandardScaler()
X_train_s = scaler.fit_transform(X_train)
X_val_s = scaler.transform(X_val)
print("train/val shape:", X_train_s.shape, X_val_s.shape)

class_weight = None
if USE_CLASS_WEIGHT:
    weight_values = compute_class_weight(
        class_weight="balanced",
        classes=np.array([0, 1]),
        y=y_train.to_numpy(),
    )
    class_weight = {0: float(weight_values[0]), 1: float(weight_values[1])}
    print("class_weight:", class_weight)


# ---- 4. MLP 기본 골격 (입력 -> 은닉 2층 -> 출력) ----
def build_model(input_dim, hidden_units=HIDDEN_UNITS):
    model = keras.Sequential([
        layers.Input(shape=(input_dim,)),
        layers.Dense(hidden_units[0], activation="relu"),
        layers.Dense(hidden_units[1], activation="relu"),
        layers.Dense(1, activation="sigmoid"),
    ])
    model.compile(
        optimizer=OPTIMIZER,
        loss="binary_crossentropy",
        metrics=[
            keras.metrics.Recall(name="recall"),
            keras.metrics.AUC(curve="PR", name="pr_auc"),
        ],
    )
    return model


model = build_model(X_train_s.shape[1])
model.summary()

# ---- 5. 학습 ----
callbacks = []
if EARLY_STOPPING:
    callbacks.append(
        keras.callbacks.EarlyStopping(
            monitor=ES_MONITOR,
            mode=ES_MODE,
            patience=ES_PATIENCE,
            restore_best_weights=RESTORE_BEST_WEIGHTS,
        )
    )

history = model.fit(
    X_train_s, y_train,
    validation_data=(X_val_s, y_val),
    epochs=MAX_EPOCHS,
    batch_size=BATCH_SIZE,
    class_weight=class_weight,
    callbacks=callbacks,
    verbose=1,
)

print("[history keys]", list(history.history.keys()))
n_epochs = len(history.history["loss"])
print("[trained epochs]", n_epochs)
if "val_pr_auc" in history.history:
    best_epoch = int(np.argmax(history.history["val_pr_auc"])) + 1
    print("[best epoch by val_pr_auc]", best_epoch)
else:
    print("[best epoch by val_pr_auc] val_pr_auc 키 없음")
print("[class_weight]", class_weight)

# ---- 6. 평가: 검증셋 Recall / F1 / PR-AUC (F1은 임계값 0.5 기준) ----
y_val_pred_prob = model.predict(X_val_s).ravel()
y_val_pred = (y_val_pred_prob >= 0.5).astype(int)

tn, fp, fn, tp = confusion_matrix(y_val, y_val_pred).ravel()
recall = recall_score(y_val, y_val_pred)
precision = precision_score(y_val, y_val_pred)
f1 = f1_score(y_val, y_val_pred)  # 임계값 0.5
pr_auc = average_precision_score(y_val, y_val_pred_prob)

print(f"[혼동행렬] TN={tn} FP={fp} FN={fn} TP={tp}")
print(
    f"\n[검증셋] Recall: {recall:.4f}  Precision: {precision:.4f}  "
    f"F1: {f1:.4f}  PR-AUC: {pr_auc:.4f}"
)
print(classification_report(y_val, y_val_pred, target_names=["No Purchase", "Purchase"]))

optimizer_name = OPTIMIZER.__class__.__name__
early_stop_log = (
    f"{ES_MONITOR}, patience={ES_PATIENCE}, mode={ES_MODE}, restore_best={RESTORE_BEST_WEIGHTS}"
    if EARLY_STOPPING
    else "없음"
)


# ---- 7. 실험 로그 엑셀에 자동 기록 ----
def log_to_excel(log_path, entry, sheet_name="실험로그"):
    """팀 공통 엑셀 양식(실험로그 시트) 맨 아래 빈 줄에 한 행 추가."""
    wb = openpyxl.load_workbook(log_path)
    ws = wb[sheet_name]

    # A열이 비어있는 첫 행 찾기 (헤더는 1행)
    next_row = 2
    while ws.cell(row=next_row, column=1).value is not None:
        next_row += 1

    exp_id = f"C1-{next_row - 1:03d}"

    row_values = [
        exp_id,
        entry["date"],
        entry["stage"],
        entry["framework"],
        entry["prep_version"],
        entry["input_dim"],
        entry["hidden_layers"],
        entry["nodes"],
        entry["optimizer"],
        entry["lr"],
        entry["batch_size"],
        entry["max_epochs"],
        entry["early_stop"],
        entry["seed"],
        entry["recall"],
        entry["f1"],
        entry["pr_auc"],
        entry["memo"],
    ]
    for col, value in enumerate(row_values, start=1):
        ws.cell(row=next_row, column=col, value=value)

    wb.save(log_path)
    print(f"[로그 저장 완료] {exp_id} -> {log_path}")
    return exp_id


log_entry = {
    "date": date.today().isoformat(),
    "stage": "baseline",
    "framework": f"tensorflow(keras) {tf.__version__}",
    "prep_version": PREP_VERSION,
    "input_dim": X_train_s.shape[1],
    "hidden_layers": len(HIDDEN_UNITS),
    "nodes": str(HIDDEN_UNITS),
    "optimizer": optimizer_name,
    "lr": LEARNING_RATE,
    "batch_size": BATCH_SIZE,
    "max_epochs": MAX_EPOCHS,
    "early_stop": early_stop_log,
    "seed": SEED,
    "recall": round(float(recall), 4),
    "f1": round(float(f1), 4),
    "pr_auc": round(float(pr_auc), 4),
    "memo": "공식 분할 기준",
}

log_to_excel(LOG_PATH, log_entry)
