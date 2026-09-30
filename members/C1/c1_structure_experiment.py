"""
C1 담당 - MLP 구조 실험 (은닉층 수 × 노드 수 × 시드)
전처리·학습 조건은 c1_mlp_baseline.py와 동일. 보류 분할은 읽지 않음.
"""

import os
from datetime import date

import numpy as np
import pandas as pd
import tensorflow as tf
from tensorflow import keras
from tensorflow.keras import layers
from sklearn.preprocessing import StandardScaler
from sklearn.utils.class_weight import compute_class_weight
from sklearn.metrics import recall_score, f1_score, average_precision_score, precision_score
import openpyxl

# ---- 실험 조건 (임시값. TODO: C2와 합의 후 확정) ----
LEARNING_RATE = 1e-3  # Adam 학습률. TODO: C2와 합의 후 확정
BATCH_SIZE = 32  # 미니배치 크기. TODO: C2와 합의 후 확정
MAX_EPOCHS = 50  # 최대 에폭 (조기종료가 더 일찍 끝낼 수 있음). TODO: C2와 합의 후 확정
EARLY_STOPPING = True  # 조기종료 사용 여부. TODO: C2와 합의 후 확정
ES_MONITOR = "val_pr_auc"  # 조기종료가 보는 지표. TODO: C2와 합의 후 확정
ES_MODE = "max"  # 지표가 커질수록 좋다고 봄. TODO: C2와 합의 후 확정
ES_PATIENCE = 5  # 개선이 없어도 기다리는 에폭 수. TODO: C2와 합의 후 확정
RESTORE_BEST_WEIGHTS = True  # 가장 좋았던 가중치로 되돌림. TODO: C2와 합의 후 확정
USE_CLASS_WEIGHT = True  # 구매/비구매 비율로 class_weight 사용. TODO: C2와 합의 후 확정
# OPTIMIZER는 clear_session 이후 run마다 Adam(learning_rate=LEARNING_RATE)로 새로 만듦. TODO: C2와 합의 후 확정

PREP_VERSION = "temp_v1(원-핫, A공통전처리 전)"
LAYER_GRID = [1, 2, 3]
UNIT_GRID = [32, 64, 128]
SEED_GRID = [42, 43, 44]
CSV_COLUMNS = [
    "layers", "units", "seed", "recall", "precision", "f1", "pr_auc",
    "epochs", "best_epoch", "n_params", "input_dim",
]

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "..", "..", "data")
FINAL_DIR = os.path.join(DATA_DIR, "final")
LOG_PATH = os.path.join(BASE_DIR, "experiment_log_C1.xlsx")
CSV_PATH = os.path.join(BASE_DIR, "structure_results.csv")


def preprocess(df, ref_columns=None):
    df = df.copy()
    df["Weekend"] = df["Weekend"].astype(int)
    y = df["Revenue"].astype(int) if "Revenue" in df.columns else None
    X = df.drop(columns=["Revenue"]) if "Revenue" in df.columns else df
    if ref_columns is None:
        X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=True)
    else:
        X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=False)
        X = X.reindex(columns=ref_columns, fill_value=0)
    return X, y


def load_data():
    train_df = pd.read_csv(os.path.join(FINAL_DIR, "train_clean.csv"))
    val_df = pd.read_csv(os.path.join(FINAL_DIR, "validation_final.csv"))
    X_train, y_train = preprocess(train_df)
    X_val, y_val = preprocess(val_df, ref_columns=X_train.columns)
    X_val = X_val[X_train.columns]
    assert list(X_train.columns) == list(X_val.columns)

    scaler = StandardScaler()
    X_train_s = scaler.fit_transform(X_train)
    X_val_s = scaler.transform(X_val)

    class_weight = None
    if USE_CLASS_WEIGHT:
        weight_values = compute_class_weight(
            class_weight="balanced",
            classes=np.array([0, 1]),
            y=y_train.to_numpy(),
        )
        class_weight = {0: float(weight_values[0]), 1: float(weight_values[1])}

    return {
        "X_train_s": X_train_s,
        "y_train": y_train,
        "X_val_s": X_val_s,
        "y_val": y_val,
        "class_weight": class_weight,
        "input_dim": X_train_s.shape[1],
    }


def build_model(n_layers, units, input_dim):
    hidden = [layers.Dense(units, activation="relu") for _ in range(n_layers)]
    model = keras.Sequential(
        [layers.Input(shape=(input_dim,))]
        + hidden
        + [layers.Dense(1, activation="sigmoid")]
    )
    model.compile(
        optimizer=keras.optimizers.Adam(learning_rate=LEARNING_RATE),
        loss="binary_crossentropy",
        metrics=[
            keras.metrics.Recall(name="recall"),
            keras.metrics.AUC(curve="PR", name="pr_auc"),
        ],
    )
    return model


def run_one(n_layers, units, seed, data):
    keras.backend.clear_session()
    keras.utils.set_random_seed(seed)

    model = build_model(n_layers, units, data["input_dim"])
    callbacks = []
    if EARLY_STOPPING:
        callbacks.append(
            keras.callbacks.EarlyStopping(
                monitor=ES_MONITOR,
                mode=ES_MODE,
                patience=ES_PATIENCE,
                restore_best_weights=RESTORE_BEST_WEIGHTS,
            )
        )

    history = model.fit(
        data["X_train_s"],
        data["y_train"],
        validation_data=(data["X_val_s"], data["y_val"]),
        epochs=MAX_EPOCHS,
        batch_size=BATCH_SIZE,
        class_weight=data["class_weight"],
        callbacks=callbacks,
        verbose=0,
    )

    n_epochs = len(history.history["loss"])
    if "val_pr_auc" in history.history:
        best_epoch = int(np.argmax(history.history["val_pr_auc"])) + 1
    else:
        best_epoch = None

    y_prob = model.predict(data["X_val_s"], verbose=0).ravel()
    y_pred = (y_prob >= 0.5).astype(int)
    y_true = data["y_val"]

    return {
        "layers": n_layers,
        "units": units,
        "seed": seed,
        "recall": float(recall_score(y_true, y_pred)),
        "precision": float(precision_score(y_true, y_pred, zero_division=0)),
        "f1": float(f1_score(y_true, y_pred)),
        "pr_auc": float(average_precision_score(y_true, y_prob)),
        "epochs": n_epochs,
        "best_epoch": best_epoch,
        "n_params": int(model.count_params()),
        "input_dim": int(data["input_dim"]),
    }


def all_combos():
    return [
        (layers_n, units, seed)
        for layers_n in LAYER_GRID
        for units in UNIT_GRID
        for seed in SEED_GRID
    ]


def load_done_keys(path):
    if not os.path.exists(path):
        return set()
    df = pd.read_csv(path)
    if df.empty:
        return set()
    return set(zip(df["layers"].astype(int), df["units"].astype(int), df["seed"].astype(int)))


def append_csv(path, row):
    df = pd.DataFrame([{k: row[k] for k in CSV_COLUMNS}])
    header = not os.path.exists(path)
    df.to_csv(path, mode="a", header=header, index=False)


def excel_existing_structure_keys(ws):
    keys = set()
    row = 2
    while ws.cell(row=row, column=1).value is not None:
        if ws.cell(row=row, column=3).value == "구조":
            layers_n = ws.cell(row=row, column=7).value
            units = ws.cell(row=row, column=8).value
            seed = ws.cell(row=row, column=14).value
            try:
                keys.add((int(layers_n), int(units), int(seed)))
            except (TypeError, ValueError):
                pass
        row += 1
    return keys, row


def write_excel_from_csv(csv_path, log_path):
    results = pd.read_csv(csv_path)
    if len(results) < 27:
        print(f"CSV가 {len(results)}행이라 엑셀 27행 추가를 건너뜀.")
        return

    early_stop_log = (
        f"{ES_MONITOR}, patience={ES_PATIENCE}, mode={ES_MODE}, restore_best={RESTORE_BEST_WEIGHTS}"
        if EARLY_STOPPING
        else "없음"
    )
    try:
        wb = openpyxl.load_workbook(log_path)
        ws = wb["실험로그"]
        existing, next_row = excel_existing_structure_keys(ws)
        added = 0
        for _, row in results.iterrows():
            key = (int(row["layers"]), int(row["units"]), int(row["seed"]))
            if key in existing:
                continue
            exp_id = f"C1-{next_row - 1:03d}"
            values = [
                exp_id,
                date.today().isoformat(),
                "구조",
                f"tensorflow(keras) {tf.__version__}",
                PREP_VERSION,
                int(row["input_dim"]),
                int(row["layers"]),
                int(row["units"]),
                "Adam",
                LEARNING_RATE,
                BATCH_SIZE,
                MAX_EPOCHS,
                early_stop_log,
                int(row["seed"]),
                round(float(row["recall"]), 4),
                round(float(row["f1"]), 4),
                round(float(row["pr_auc"]), 4),
                "공식 분할 기준",
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=next_row, column=col, value=value)
            next_row += 1
            added += 1
        wb.save(log_path)
        print(f"[엑셀] 실험로그에 {added}행 추가 (이미 있던 구조 행은 건너뜀)")
    except PermissionError:
        print("엑셀을 닫고 다시 실행하세요")


def print_summary(csv_path):
    df = pd.read_csv(csv_path)
    grouped = (
        df.groupby(["layers", "units"], as_index=False)
        .agg(
            n=("seed", "count"),
            pr_auc_mean=("pr_auc", "mean"),
            pr_auc_std=("pr_auc", "std"),
            recall_mean=("recall", "mean"),
            recall_std=("recall", "std"),
            f1_mean=("f1", "mean"),
            f1_std=("f1", "std"),
            epochs_mean=("epochs", "mean"),
        )
        .sort_values("pr_auc_mean", ascending=False)
    )
    print("\n[조합별 평균±표준편차] (표본 표준편차 n-1, PR-AUC 평균 내림차순)")
    print(
        f"{'layers':>6} {'units':>5} {'n':>3} "
        f"{'PR-AUC':>18} {'Recall':>18} {'F1':>18} {'epochs평균':>10}"
    )
    for _, r in grouped.iterrows():
        def fmt(mean, std):
            if pd.isna(std):
                return f"{mean:.4f}±NA"
            return f"{mean:.4f}±{std:.4f}"

        print(
            f"{int(r['layers']):>6} {int(r['units']):>5} {int(r['n']):>3} "
            f"{fmt(r['pr_auc_mean'], r['pr_auc_std']):>18} "
            f"{fmt(r['recall_mean'], r['recall_std']):>18} "
            f"{fmt(r['f1_mean'], r['f1_std']):>18} "
            f"{r['epochs_mean']:>10.1f}"
        )
    return grouped, df


def main(smoke=False):
    data = load_data()
    combos = all_combos()
    done = load_done_keys(CSV_PATH)
    target = [(2, 64, 42)] if smoke else combos

    for layers_n, units, seed in target:
        k = combos.index((layers_n, units, seed)) + 1
        if (layers_n, units, seed) in done:
            print(f"[{k}/27] layers={layers_n} units={units} seed={seed} SKIP (CSV에 있음)")
            continue
        row = run_one(layers_n, units, seed, data)
        append_csv(CSV_PATH, row)
        print(
            f"[{k}/27] layers={layers_n} units={units} seed={seed} "
            f"epochs={row['epochs']} best={row['best_epoch']} "
            f"Recall={row['recall']:.4f} Prec={row['precision']:.4f} "
            f"F1={row['f1']:.4f} PR-AUC={row['pr_auc']:.4f}"
        )

    if smoke:
        return

    done = load_done_keys(CSV_PATH)
    if len(done) == 27:
        write_excel_from_csv(CSV_PATH, LOG_PATH)
        print(
            "집계 시트는 수식(단계='구조', 은닉층·노드 숫자 매칭)으로 자동 계산되므로 "
            "손대지 않음."
        )
        print_summary(CSV_PATH)
    else:
        print(f"완료된 조합 {len(done)}/27. 엑셀 추가는 27회가 끝난 뒤에만 함.")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--smoke",
        action="store_true",
        help="(2, 64, 42) 한 조합만 실행",
    )
    args = parser.parse_args()
    main(smoke=args.smoke)
