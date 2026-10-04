"""
C1 원핫 인코딩 범위 비교 (조건 A/B/C × 시드 3).
모델·학습·평가는 c1_structure_experiment.py를 import해서 재사용한다.
기존 구조/baseline 파일은 수정하지 않는다.
"""

from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import tensorflow as tf
from sklearn.preprocessing import StandardScaler
from sklearn.utils.class_weight import compute_class_weight
import openpyxl

C1_DIR = Path(__file__).resolve().parent
if str(C1_DIR) not in sys.path:
    sys.path.insert(0, str(C1_DIR))

import c1_structure_experiment as se

PROJECT_ROOT = C1_DIR.parent.parent
FINAL_DIR = PROJECT_ROOT / "data" / "final"
LOG_PATH = C1_DIR / "experiment_log_C1.xlsx"
CSV_PATH = C1_DIR / "onehot_results.csv"

# 구조 실험 격자에서 layers=2, units=64 → 은닉층 Dense(64) 두 개 (64, 64)
N_LAYERS = 2
UNITS = 64
SEEDS = [42, 43, 44]

CONDITIONS = {
    "A": ["Month", "VisitorType"],
    "B": ["Month", "VisitorType", "TrafficType"],
    "C": ["Month", "VisitorType", "TrafficType", "OperatingSystems", "Browser", "Region"],
}

CSV_COLUMNS = [
    "조건", "시드", "입력 차원", "실제 에폭", "best 에폭",
    "val Recall", "F1", "PR-AUC",
]


def preprocess(df, onehot_cols, ref_columns=None):
    """C1과 동일: 학습은 drop_first=True, 검증은 drop_first=False 후 reindex.

    Month/VisitorType은 기존처럼 drop_first=True.
    추가 범주형(TrafficType 등)은 정수 코드의 모든 수준을 남겨
    입력 차원이 A=26, B=44, C=71이 되게 한다.
    """
    df = df.copy()
    df["Weekend"] = df["Weekend"].astype(int)
    y = df["Revenue"].astype(int) if "Revenue" in df.columns else None
    X = df.drop(columns=["Revenue"]) if "Revenue" in df.columns else df

    base_cols = [c for c in onehot_cols if c in ("Month", "VisitorType")]
    extra_cols = [c for c in onehot_cols if c not in ("Month", "VisitorType")]
    for col in extra_cols:
        X[col] = X[col].astype(str)

    if ref_columns is None:
        if base_cols:
            X = pd.get_dummies(X, columns=base_cols, drop_first=True)
        if extra_cols:
            X = pd.get_dummies(X, columns=extra_cols, drop_first=False)
    else:
        dummy_cols = base_cols + extra_cols
        if dummy_cols:
            X = pd.get_dummies(X, columns=dummy_cols, drop_first=False)
        X = X.reindex(columns=ref_columns, fill_value=0)
    return X, y


def load_data(onehot_cols):
    train_df = pd.read_csv(FINAL_DIR / "train_clean.csv")
    val_df = pd.read_csv(FINAL_DIR / "validation_final.csv")
    X_train, y_train = preprocess(train_df, onehot_cols)
    X_val, y_val = preprocess(val_df, onehot_cols, ref_columns=X_train.columns)
    X_val = X_val[X_train.columns]
    assert list(X_train.columns) == list(X_val.columns)

    scaler = StandardScaler()
    X_train_s = scaler.fit_transform(X_train)
    X_val_s = scaler.transform(X_val)

    class_weight = None
    if se.USE_CLASS_WEIGHT:
        weight_values = compute_class_weight(
            class_weight="balanced",
            classes=np.array([0, 1]),
            y=y_train.to_numpy(),
        )
        class_weight = {0: float(weight_values[0]), 1: float(weight_values[1])}

    return {
        "X_train_s": X_train_s,
        "y_train": y_train,
        "X_val_s": X_val_s,
        "y_val": y_val,
        "class_weight": class_weight,
        "input_dim": int(X_train_s.shape[1]),
    }


def load_done_keys(path):
    if not path.exists():
        return set()
    df = pd.read_csv(path)
    if df.empty:
        return set()
    return set(zip(df["조건"].astype(str), df["시드"].astype(int)))


def append_csv(row):
    pd.DataFrame([{k: row[k] for k in CSV_COLUMNS}]).to_csv(
        CSV_PATH, mode="a", header=not CSV_PATH.exists(), index=False
    )


def append_excel(row, onehot_cols, input_dim):
    early_stop_log = (
        f"{se.ES_MONITOR}, patience={se.ES_PATIENCE}, mode={se.ES_MODE}, "
        f"restore_best={se.RESTORE_BEST_WEIGHTS}"
        if se.EARLY_STOPPING
        else "없음"
    )
    memo = (
        f"공식 분할 기준, 조건 {row['조건']}, "
        f"원핫={onehot_cols}, 입력차원={input_dim}"
    )
    try:
        wb = openpyxl.load_workbook(LOG_PATH)
        ws = wb["실험로그"]
        next_row = 2
        while ws.cell(row=next_row, column=1).value is not None:
            next_row += 1
        exp_id = f"C1-{next_row - 1:03d}"
        values = [
            exp_id,
            date.today().isoformat(),
            "원핫비교",
            f"tensorflow(keras) {tf.__version__}",
            se.PREP_VERSION,
            input_dim,
            N_LAYERS,
            UNITS,
            "Adam",
            se.LEARNING_RATE,
            se.BATCH_SIZE,
            se.MAX_EPOCHS,
            early_stop_log,
            int(row["시드"]),
            round(float(row["val Recall"]), 4),
            round(float(row["F1"]), 4),
            round(float(row["PR-AUC"]), 4),
            memo,
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=next_row, column=col, value=value)
        wb.save(LOG_PATH)
        print(f"[엑셀] {exp_id}")
    except PermissionError:
        print("엑셀을 닫고 다시 실행하세요")


def print_summary():
    df = pd.read_csv(CSV_PATH)
    print("\n[입력 차원]")
    print(df.groupby("조건")["입력 차원"].first().to_string())
    grouped = (
        df.groupby("조건")
        .agg(
            n=("시드", "count"),
            pr_mean=("PR-AUC", "mean"),
            pr_std=("PR-AUC", "std"),
            f1_mean=("F1", "mean"),
            f1_std=("F1", "std"),
            rec_mean=("val Recall", "mean"),
            rec_std=("val Recall", "std"),
            ep_mean=("실제 에폭", "mean"),
        )
        .reindex(["A", "B", "C"])
    )
    print("\n[조건별 평균±표준편차 ddof=1]")
    for name, r in grouped.iterrows():
        print(
            f"  {name}: PR-AUC {r['pr_mean']:.4f}±{r['pr_std']:.4f}  "
            f"F1 {r['f1_mean']:.4f}±{r['f1_std']:.4f}  "
            f"Recall {r['rec_mean']:.4f}±{r['rec_std']:.4f}  "
            f"에폭평균 {r['ep_mean']:.1f}  n={int(r['n'])}"
        )


def main(smoke=False):
    seeds = [42] if smoke else SEEDS
    saved_epochs = se.MAX_EPOCHS
    if smoke:
        se.MAX_EPOCHS = 2

    try:
        for cond_name, onehot_cols in CONDITIONS.items():
            data = load_data(onehot_cols)
            print(f"[dim] {cond_name} {onehot_cols} -> {data['input_dim']}")
            for seed in seeds:
                if not smoke and (cond_name, seed) in load_done_keys(CSV_PATH):
                    print(f"SKIP {cond_name} seed={seed}")
                    continue
                result = se.run_one(N_LAYERS, UNITS, seed, data)
                row = {
                    "조건": cond_name,
                    "시드": seed,
                    "입력 차원": result["input_dim"],
                    "실제 에폭": result["epochs"],
                    "best 에폭": result["best_epoch"],
                    "val Recall": result["recall"],
                    "F1": result["f1"],
                    "PR-AUC": result["pr_auc"],
                }
                print(
                    f"[{cond_name} seed={seed}] dim={result['input_dim']} "
                    f"epochs={result['epochs']} best={result['best_epoch']} "
                    f"Recall={result['recall']:.4f} F1={result['f1']:.4f} "
                    f"PR-AUC={result['pr_auc']:.4f}"
                )
                if not smoke:
                    append_csv(row)
                    append_excel(row, onehot_cols, result["input_dim"])
    finally:
        se.MAX_EPOCHS = saved_epochs

    if not smoke and CSV_PATH.exists():
        print_summary()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--smoke", action="store_true")
    args = parser.parse_args()
    main(smoke=args.smoke)
