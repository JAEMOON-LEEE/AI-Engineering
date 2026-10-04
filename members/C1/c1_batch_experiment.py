"""
C1 담당 - MLP 배치 크기 실험 (batch × 시드)
옵티마이저·학습률은 adam / 1e-3, 구조는 2층×64로 고정한다.
전처리·시드 고정·평가는 c1_optimizer_experiment.py와 동일. 보류 분할은 읽지 않음.
"""

import math
import os
import sys
import time
from datetime import date

import numpy as np
import pandas as pd
import tensorflow as tf
from tensorflow import keras
from tensorflow.keras import layers
from sklearn.preprocessing import StandardScaler
from sklearn.utils.class_weight import compute_class_weight
from sklearn.metrics import recall_score, f1_score, average_precision_score, precision_score
import openpyxl

# ---- 실험 조건 (C1 기준 확정(C2 동의)) ----
LEARNING_RATE = 1e-3  # Adam 학습률. 옵티마이저 실험 결과로 확정. C1 기준 확정(C2 동의)
MAX_EPOCHS = 50  # 최대 에폭 (조기종료가 더 일찍 끝낼 수 있음). C1 기준 확정(C2 동의)
EARLY_STOPPING = True  # 조기종료 사용 여부. C1 기준 확정(C2 동의)
ES_MONITOR = "val_pr_auc"  # 조기종료가 보는 지표. C1 기준 확정(C2 동의)
ES_MODE = "max"  # 지표가 커질수록 좋다고 봄. C1 기준 확정(C2 동의)
ES_PATIENCE = 5  # 개선이 없어도 기다리는 에폭 수. C1 기준 확정(C2 동의)
RESTORE_BEST_WEIGHTS = True  # 가장 좋았던 가중치로 되돌림. C1 기준 확정(C2 동의)
USE_CLASS_WEIGHT = True  # 구매/비구매 비율로 class_weight 사용. C1 기준 확정(C2 동의)
# OPTIMIZER는 clear_session 이후 run마다 make_optimizer로 새로 만듦. C1 기준 확정(C2 동의)

SGD_MOMENTUM = 0.9
ADAMW_WEIGHT_DECAY = 0.004
N_LAYERS = 2
UNITS = 64
OPTIMIZER_NAME = "adam"
BATCH_GRID = [32, 64, 128, 256]
SEED_GRID = [42, 43, 44]

PREP_VERSION = "official_v1(data/final, Month·VisitorType 원핫, 26차원)"
CSV_COLUMNS = [
    "optimizer", "lr", "batch_size", "layers", "units", "seed",
    "recall", "precision", "f1", "pr_auc",
    "epochs", "best_epoch", "hit_max_epoch", "steps_per_epoch",
    "train_seconds", "n_params", "input_dim",
]

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "..", "..", "data")
FINAL_DIR = os.path.join(DATA_DIR, "final")
LOG_PATH = os.path.join(BASE_DIR, "experiment_log_C1.xlsx")
CSV_PATH = os.path.join(BASE_DIR, "batch_results.csv")
OPTIMIZER_CSV = os.path.join(BASE_DIR, "optimizer_results.csv")


def preprocess(df, ref_columns=None):
    df = df.copy()
    df["Weekend"] = df["Weekend"].astype(int)
    y = df["Revenue"].astype(int) if "Revenue" in df.columns else None
    X = df.drop(columns=["Revenue"]) if "Revenue" in df.columns else df
    if ref_columns is None:
        X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=True)
    else:
        X = pd.get_dummies(X, columns=["Month", "VisitorType"], drop_first=False)
        X = X.reindex(columns=ref_columns, fill_value=0)
    return X, y


def load_data():
    train_df = pd.read_csv(os.path.join(FINAL_DIR, "train_clean.csv"))
    val_df = pd.read_csv(os.path.join(FINAL_DIR, "validation_final.csv"))
    X_train, y_train = preprocess(train_df)
    X_val, y_val = preprocess(val_df, ref_columns=X_train.columns)
    X_val = X_val[X_train.columns]
    assert list(X_train.columns) == list(X_val.columns)

    scaler = StandardScaler()
    X_train_s = scaler.fit_transform(X_train)
    X_val_s = scaler.transform(X_val)

    class_weight = None
    if USE_CLASS_WEIGHT:
        weight_values = compute_class_weight(
            class_weight="balanced",
            classes=np.array([0, 1]),
            y=y_train.to_numpy(),
        )
        class_weight = {0: float(weight_values[0]), 1: float(weight_values[1])}

    return {
        "X_train_s": X_train_s,
        "y_train": y_train,
        "X_val_s": X_val_s,
        "y_val": y_val,
        "class_weight": class_weight,
        "input_dim": X_train_s.shape[1],
        "n_train": int(len(y_train)),
    }


def make_optimizer(name, lr):
    if name == "sgd_momentum":
        return keras.optimizers.SGD(learning_rate=lr, momentum=SGD_MOMENTUM)
    if name == "adam":
        return keras.optimizers.Adam(learning_rate=lr)
    if name == "adamw":
        return keras.optimizers.AdamW(learning_rate=lr, weight_decay=ADAMW_WEIGHT_DECAY)
    raise ValueError(f"지원하지 않는 옵티마이저: {name}")


def build_model(n_layers, units, input_dim, optimizer_name, lr):
    hidden = [layers.Dense(units, activation="relu") for _ in range(n_layers)]
    model = keras.Sequential(
        [layers.Input(shape=(input_dim,))]
        + hidden
        + [layers.Dense(1, activation="sigmoid")]
    )
    model.compile(
        optimizer=make_optimizer(optimizer_name, lr),
        loss="binary_crossentropy",
        metrics=[
            keras.metrics.Recall(name="recall"),
            keras.metrics.AUC(curve="PR", name="pr_auc"),
        ],
    )
    return model


def run_one(batch_size, seed, data):
    keras.backend.clear_session()
    keras.utils.set_random_seed(seed)

    model = build_model(N_LAYERS, UNITS, data["input_dim"], OPTIMIZER_NAME, LEARNING_RATE)
    callbacks = []
    if EARLY_STOPPING:
        callbacks.append(
            keras.callbacks.EarlyStopping(
                monitor=ES_MONITOR,
                mode=ES_MODE,
                patience=ES_PATIENCE,
                restore_best_weights=RESTORE_BEST_WEIGHTS,
            )
        )

    t0 = time.perf_counter()
    history = model.fit(
        data["X_train_s"],
        data["y_train"],
        validation_data=(data["X_val_s"], data["y_val"]),
        epochs=MAX_EPOCHS,
        batch_size=batch_size,
        class_weight=data["class_weight"],
        callbacks=callbacks,
        verbose=0,
    )
    train_seconds = time.perf_counter() - t0

    n_epochs = len(history.history["loss"])
    if "val_pr_auc" in history.history:
        best_epoch = int(np.argmax(history.history["val_pr_auc"])) + 1
    else:
        best_epoch = None

    y_prob = model.predict(data["X_val_s"], verbose=0).ravel()
    y_pred = (y_prob >= 0.5).astype(int)
    y_true = data["y_val"]
    steps_per_epoch = int(math.ceil(data["n_train"] / batch_size))

    return {
        "optimizer": OPTIMIZER_NAME,
        "lr": LEARNING_RATE,
        "batch_size": int(batch_size),
        "layers": N_LAYERS,
        "units": UNITS,
        "seed": seed,
        "recall": float(recall_score(y_true, y_pred)),
        "precision": float(precision_score(y_true, y_pred, zero_division=0)),
        "f1": float(f1_score(y_true, y_pred)),
        "pr_auc": float(average_precision_score(y_true, y_prob)),
        "epochs": n_epochs,
        "best_epoch": best_epoch,
        "hit_max_epoch": bool(n_epochs == MAX_EPOCHS),
        "steps_per_epoch": steps_per_epoch,
        "train_seconds": float(train_seconds),
        "n_params": int(model.count_params()),
        "input_dim": int(data["input_dim"]),
    }


def all_combos():
    return [
        (batch_size, seed)
        for batch_size in BATCH_GRID
        for seed in SEED_GRID
    ]


def combo_key(batch_size, seed):
    return (int(batch_size), int(seed))


def load_done_keys(path):
    if not os.path.exists(path):
        return set()
    df = pd.read_csv(path)
    if df.empty:
        return set()
    return set(combo_key(b, s) for b, s in zip(df["batch_size"], df["seed"]))


def append_csv(path, row):
    df = pd.DataFrame([{k: row[k] for k in CSV_COLUMNS}])
    header = not os.path.exists(path)
    df.to_csv(path, mode="a", header=header, index=False)


def excel_existing_batch_keys(ws):
    keys = set()
    row = 2
    while ws.cell(row=row, column=1).value is not None:
        if ws.cell(row=row, column=3).value == "배치":
            batch_size = ws.cell(row=row, column=11).value
            seed = ws.cell(row=row, column=14).value
            try:
                keys.add((int(batch_size), int(seed)))
            except (TypeError, ValueError):
                pass
        row += 1
    return keys, row


def write_excel_from_csv(csv_path, log_path):
    results = pd.read_csv(csv_path)
    if len(results) < 12:
        print(f"CSV가 {len(results)}행이라 엑셀 12행 추가를 건너뜀.")
        return

    early_stop_log = (
        f"{ES_MONITOR}, patience={ES_PATIENCE}, mode={ES_MODE}, restore_best={RESTORE_BEST_WEIGHTS}"
        if EARLY_STOPPING
        else "없음"
    )
    try:
        wb = openpyxl.load_workbook(log_path)
        ws = wb["실험로그"]
        existing, next_row = excel_existing_batch_keys(ws)
        added = 0
        first_id = None
        last_id = None
        for _, row in results.iterrows():
            key = (int(row["batch_size"]), int(row["seed"]))
            if key in existing:
                continue
            exp_id = f"C1-{next_row - 1:03d}"
            if first_id is None:
                first_id = exp_id
            last_id = exp_id
            values = [
                exp_id,
                date.today().isoformat(),
                "배치",
                f"tensorflow(keras) {tf.__version__}",
                PREP_VERSION,
                int(row["input_dim"]),
                int(row["layers"]),
                int(row["units"]),
                "Adam",
                LEARNING_RATE,
                int(row["batch_size"]),
                MAX_EPOCHS,
                early_stop_log,
                int(row["seed"]),
                round(float(row["recall"]), 4),
                round(float(row["f1"]), 4),
                round(float(row["pr_auc"]), 4),
                "공식 분할 기준",
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=next_row, column=col, value=value)
            next_row += 1
            added += 1
        wb.save(log_path)
        if added:
            print(f"[엑셀] 실험로그에 {added}행 추가 {first_id}~{last_id} (이미 있던 배치 행은 건너뜀)")
        else:
            print("[엑셀] 추가할 배치 행 없음 (이미 있음)")
    except PermissionError:
        print("엑셀을 닫고 다시 실행하세요")
        sys.exit(1)


def print_summary(csv_path):
    df = pd.read_csv(csv_path)
    grouped = (
        df.groupby(["batch_size"], as_index=False)
        .agg(
            n=("seed", "count"),
            pr_auc_mean=("pr_auc", "mean"),
            pr_auc_std=("pr_auc", "std"),
            recall_mean=("recall", "mean"),
            recall_std=("recall", "std"),
            f1_mean=("f1", "mean"),
            f1_std=("f1", "std"),
            epochs_mean=("epochs", "mean"),
            best_epoch_mean=("best_epoch", "mean"),
            steps=("steps_per_epoch", "first"),
            hit_max=("hit_max_epoch", "sum"),
            train_sec_mean=("train_seconds", "mean"),
        )
        .sort_values("pr_auc_mean", ascending=False)
    )
    print("\n[배치별 평균±표준편차] (표본 표준편차 n-1, PR-AUC 평균 내림차순)")
    print(
        f"{'batch':>6} {'n':>3} "
        f"{'PR-AUC':>18} {'Recall':>18} {'F1':>18} "
        f"{'epochs평균':>10} {'best평균':>8} {'steps':>6} {'hit_max':>8} {'초/회':>8}"
    )

    def fmt(mean, std):
        if pd.isna(std):
            return f"{mean:.4f}±NA"
        return f"{mean:.4f}±{std:.4f}"

    for _, r in grouped.iterrows():
        print(
            f"{int(r['batch_size']):>6} {int(r['n']):>3} "
            f"{fmt(r['pr_auc_mean'], r['pr_auc_std']):>18} "
            f"{fmt(r['recall_mean'], r['recall_std']):>18} "
            f"{fmt(r['f1_mean'], r['f1_std']):>18} "
            f"{r['epochs_mean']:>10.1f} {r['best_epoch_mean']:>8.1f} "
            f"{int(r['steps']):>6} {int(r['hit_max']):>8} {r['train_sec_mean']:>8.1f}"
        )
    return grouped, df


def compare_smoke_to_optimizer(row):
    opt = pd.read_csv(OPTIMIZER_CSV)
    ref = opt[
        (opt["optimizer"] == "adam")
        & (np.isclose(opt["lr"].astype(float), 1e-3))
        & (opt["seed"] == 42)
    ]
    if ref.empty:
        print("[재현성] optimizer_results.csv에 adam/1e-3/seed=42 행이 없음")
        return False
    ref = ref.iloc[0]
    print("[재현성] batch 32 seed=42 vs adam 1e-3 seed=42")
    ok = True
    for field in ["recall", "f1", "pr_auc", "epochs"]:
        a = int(row[field]) if field == "epochs" else float(row[field])
        b = int(ref[field]) if field == "epochs" else float(ref[field])
        status = "일치" if (a == b if field == "epochs" else np.isclose(a, b, rtol=1e-12, atol=1e-12)) else "불일치"
        if status == "불일치":
            ok = False
        print(f"  {field}: batch={a} optimizer={b} -> {status}")
    return ok


def main(smoke=False):
    data = load_data()
    combos = all_combos()
    done = load_done_keys(CSV_PATH)
    target = [(32, 42)] if smoke else combos

    last_row = None
    for batch_size, seed in target:
        k = combos.index((batch_size, seed)) + 1
        key = combo_key(batch_size, seed)
        if key in done:
            print(f"[{k}/12] batch={batch_size} seed={seed} SKIP (CSV에 있음)")
            if smoke and os.path.exists(CSV_PATH):
                prev = pd.read_csv(CSV_PATH)
                hit = prev[(prev["batch_size"] == 32) & (prev["seed"] == 42)]
                if not hit.empty:
                    last_row = hit.iloc[-1].to_dict()
            continue
        row = run_one(batch_size, seed, data)
        append_csv(CSV_PATH, row)
        last_row = row
        print(
            f"[{k}/12] batch={batch_size} seed={seed} steps={row['steps_per_epoch']} "
            f"epochs={row['epochs']} best={row['best_epoch']} hit_max={row['hit_max_epoch']} "
            f"sec={row['train_seconds']:.1f} "
            f"Recall={row['recall']:.4f} Prec={row['precision']:.4f} "
            f"F1={row['f1']:.4f} PR-AUC={row['pr_auc']:.4f}"
        )

    if smoke:
        if last_row is None:
            print("[재현성] 스모크 결과를 찾지 못함")
            sys.exit(1)
        if not compare_smoke_to_optimizer(last_row):
            print("[재현성] 불일치. 여기서 중단.")
            sys.exit(1)
        print("[재현성] 일치. 전체 실험 진행 가능.")
        return

    done = load_done_keys(CSV_PATH)
    if len(done) == 12:
        write_excel_from_csv(CSV_PATH, LOG_PATH)
        print("집계 시트는 건드리지 않음.")
        print_summary(CSV_PATH)
    else:
        print(f"완료된 조합 {len(done)}/12. 엑셀 추가는 12회가 끝난 뒤에만 함.")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--smoke",
        action="store_true",
        help="(batch=32, seed=42) 한 조합만 실행 후 옵티마이저 실험과 비교",
    )
    args = parser.parse_args()
    main(smoke=args.smoke)
